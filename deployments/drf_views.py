import datetime
from collections import defaultdict
from datetime import date

from django.contrib.postgres.aggregates.general import ArrayAgg
from django.contrib.postgres.fields import ArrayField
from django.db import models
from django.db.models import Count, Q
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.utils.translation import override as translation_override
from django_filters import rest_framework as filters
from drf_spectacular.utils import extend_schema
from openpyxl import Workbook
from rest_framework import viewsets
from rest_framework.authentication import TokenAuthentication
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from reversion.views import RevisionMixin

from api.models import Country, Event, Region
from api.utils import bad_request
from api.view_filters import ListFilter
from api.visibility_class import ReadOnlyVisibilityViewsetMixin
from deployments.permissions import ERUReadinessPermission
from main.permissions import DenyGuestUserPermission
from main.serializers import CsvListMixin
from main.utils import is_tableau

from .filters import EmergencyProjectFilter, ERUOwnerFilter, ProjectFilter
from .models import (
    ERU,
    EmergencyProject,
    EmergencyProjectActivityAction,
    EmergencyProjectActivitySector,
    ERUOwner,
    ERUReadiness,
    ERUReadinessType,
    ERUType,
    OperationTypes,
    PartnerSocietyDeployment,
    Personnel,
    PersonnelDeployment,
    ProgrammeTypes,
    Project,
    RegionalProject,
    Sector,
    Statuses,
)
from .serializers import (
    AggregateDeploymentsSerializer,
    AggregatedERUAndRapidResponseSerializer,
    DeploymentByNSSerializer,
    DeploymentsByMonthSerializer,
    EmergencyProjectOptionsSerializer,
    EmergencyProjectSerializer,
    ERUOwnerMiniSerializer,
    ERUOwnerSerializer,
    ERUReadinessSerializer,
    ERUSerializer,
    GlobalProjectNSOngoingProjectsStatsSerializer,
    GlobalProjectOverviewSerializer,
    MiniERUReadinessTypeSerializer,
    PartnerDeploymentSerializer,
    PartnerDeploymentTableauSerializer,
    PersonnelCsvSerializer,
    PersonnelCsvSerializerAnon,
    PersonnelCsvSerializerSuper,
    PersonnelDeploymentSerializer,
    PersonnelSerializer,
    ProjectCsvSerializer,
    ProjectRegionMovementActivitiesSerializer,
    ProjectRegionOverviewSerializer,
    ProjectSerializer,
    RegionalProjectSerializer,
)
from .utils import get_previous_months


class ERUOwnerViewset(viewsets.ReadOnlyModelViewSet):
    authentication_classes = (TokenAuthentication,)
    # Also unauthenticated users should reach Surge page content. 2021.09.28:
    # permission_classes = (IsAuthenticated,)
    queryset = ERUOwner.objects.all()
    serializer_class = ERUOwnerSerializer
    ordering_fields = (
        "created_at",
        "updated_at",
    )
    filterset_class = ERUOwnerFilter
    search_fields = ("national_society_country__name",)  # for /docs

    @extend_schema(
        request=None,
        responses=ERUOwnerMiniSerializer(many=True),
    )
    @action(
        detail=False,
        methods=("get",),
        url_path="mini",
    )
    def mini(self, request):
        queryset = ERUOwner.objects.select_related("national_society_country").all()
        serializer = ERUOwnerMiniSerializer(queryset, many=True)
        page = self.paginate_queryset(queryset=queryset)
        if page is not None:
            serializer = ERUOwnerMiniSerializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        return Response(serializer.data)


class ERUFilter(filters.FilterSet):
    deployed_to__isnull = filters.BooleanFilter(field_name="deployed_to", lookup_expr="isnull")
    deployed_to__in = ListFilter(field_name="deployed_to__id")
    type = filters.NumberFilter(field_name="type", lookup_expr="exact")
    event = filters.NumberFilter(field_name="event", lookup_expr="exact")
    event__in = ListFilter(field_name="event")
    disaster_type = filters.NumberFilter(field_name="event__dtype", lookup_expr="exact")

    class Meta:
        model = ERU
        fields = {
            "available": ("exact",),
            "start_date": ("exact", "gt", "gte", "lt", "lte"),
            "end_date": ("exact", "gt", "gte", "lt", "lte"),
        }


class ERUViewset(viewsets.ReadOnlyModelViewSet):
    authentication_classes = (TokenAuthentication,)
    # Some figures are shown on the home page also, and not only authenticated users should see them.
    # permission_classes = (IsAuthenticated,)
    queryset = ERU.objects.select_related("eru_owner").prefetch_related(
        "deployed_to",
        "event",
        "event__appeals",
        "event__dtype",
        "event__countries",
        "event__field_reports",
        "event__field_reports__countries",
        "event__field_reports__contacts",
        "eru_owner__national_society_country",
        "eru_owner__eru_set",
        "eru_owner__eru_set__deployed_to",
    )
    # ERUSerializer uses ERUOwnerSerializer which uses ERUSetSerializer (~circle)
    serializer_class = ERUSerializer
    filterset_class = ERUFilter
    ordering_fields = (
        "type",
        "units",
        "equipment_units",
        "deployed_to__society_name",
        "event__name",
        "eru_owner__national_society_country__society_name",
        "available",
    )


class AggregatedERUAndRapidResponseViewSet(viewsets.ReadOnlyModelViewSet):
    serializer_class = AggregatedERUAndRapidResponseSerializer

    def get_queryset(self):
        today = timezone.now().date().strftime("%Y-%m-%d")
        active_personnel_prefetch = models.Prefetch(
            "personneldeployment_set__personnel_set",
            queryset=(
                Personnel.objects.filter(
                    type=Personnel.TypeChoices.RR,
                    start_date__date__lte=today,
                    end_date__date__gte=today,
                    is_active=True,
                ).select_related("country_from")
            ),
        )
        active_eru_prefetch = models.Prefetch(
            "eru_set",
            queryset=(
                ERU.objects.filter(
                    deployed_to__isnull=False,
                    start_date__date__lte=today,
                    end_date__date__gte=today,
                ).select_related(
                    "eru_owner__national_society_country",
                )
            ),
        )
        queryset = (
            Event.objects.prefetch_related(
                active_personnel_prefetch,
                active_eru_prefetch,
                "appeals",
            )
            .annotate(
                deployed_eru_count=Count(
                    "eru",
                    filter=Q(
                        eru__deployed_to__isnull=False,
                        eru__start_date__date__lte=today,
                        eru__end_date__date__gte=today,
                    ),
                    distinct=True,
                ),
                deployed_personnel_count=Count(
                    "personneldeployment__personnel",
                    filter=Q(
                        personneldeployment__personnel__type=Personnel.TypeChoices.RR,
                        personneldeployment__personnel__is_active=True,
                        personneldeployment__personnel__start_date__date__lte=today,
                        personneldeployment__personnel__end_date__date__gte=today,
                    ),
                    distinct=True,
                ),
            )
            .exclude(Q(deployed_eru_count=0) & Q(deployed_personnel_count=0))
            .order_by("-disaster_start_date")
        )
        return queryset


class PersonnelDeploymentFilter(filters.FilterSet):
    country_deployed_to = filters.NumberFilter(field_name="country_deployed_to", lookup_expr="exact")
    region_deployed_to = filters.NumberFilter(field_name="region_deployed_to", lookup_expr="exact")
    event_deployed_to = filters.NumberFilter(field_name="event_deployed_to", lookup_expr="exact")

    class Meta:
        model = PersonnelDeployment
        fields = (
            "country_deployed_to",
            "region_deployed_to",
            "event_deployed_to",
        )


class PersonnelDeploymentViewset(viewsets.ReadOnlyModelViewSet):
    authentication_classes = (TokenAuthentication,)
    permission_classes = (IsAuthenticated, DenyGuestUserPermission)
    queryset = PersonnelDeployment.objects.all()
    serializer_class = PersonnelDeploymentSerializer
    filterset_class = PersonnelDeploymentFilter
    ordering_fields = (
        "country_deployed_to",
        "region_deployed_to",
        "event_deployed_to",
    )
    search_fields = ("country_deployed_to__name", "region_deployed_to__label", "event_deployed_to__name")  #


class PersonnelFilter(filters.FilterSet):
    country_from = filters.NumberFilter(field_name="country_from", lookup_expr="exact")
    country_to = filters.NumberFilter(field_name="country_to", lookup_expr="exact")
    type = filters.CharFilter(field_name="type", lookup_expr="exact")
    event_deployed_to = filters.NumberFilter(field_name="deployment__event_deployed_to", lookup_expr="exact")
    is_active = filters.BooleanFilter(field_name="is_active", lookup_expr="exact")
    dtype = filters.NumberFilter(field_name="deployment__event_deployed_to__dtype", lookup_expr="exact")

    class Meta:
        model = Personnel
        fields = {
            "start_date": ("exact", "gt", "gte", "lt", "lte"),
            "end_date": ("exact", "gt", "gte", "lt", "lte"),
            "deployment__updated_at": ("exact", "gt", "gte", "lt", "lte"),
            "role": ("exact",),
        }


class PersonnelViewset(viewsets.ReadOnlyModelViewSet):
    authentication_classes = (TokenAuthentication,)
    # Some figures are shown on the home page also, and not only authenticated users should see them.
    # permission_classes = (IsAuthenticated,)
    queryset = Personnel.objects.all()
    filterset_class = PersonnelFilter
    ordering_fields = (
        "start_date",
        "end_date",
        "name",
        "role",
        "type",
        "country_from",
        "deployment",
    )
    search_fields = (
        "name",
        "role",
        "type",
    )  # for /docs

    def get_queryset(self):
        qs = super().get_queryset()
        return (
            qs.filter(is_active=True)
            .select_related(
                "deployment__country_deployed_to",
                "deployment__event_deployed_to",
                "deployment__event_deployed_to__dtype",
                "country_from",
                "country_to",
            )
            .prefetch_related(
                "deployment__event_deployed_to__countries",
                "deployment__event_deployed_to__appeals",
                "molnix_tags",
                "molnix_tags__groups",
            )
            .all()
        )

    def get_serializer_class(self):
        request_format_type = self.request.GET.get("format", "json")
        if request_format_type == "csv":
            if self.request.user.is_anonymous:
                return PersonnelCsvSerializerAnon
            elif self.request.user.is_superuser:
                return PersonnelCsvSerializerSuper
            return PersonnelCsvSerializer
        # if self.request.user.is_anonymous:
        #     return PersonnelSerializerAnon
        # elif self.request.user.is_superuser:
        #     return PersonnelSerializerSuper
        return PersonnelSerializer

    def get_renderer_context(self):
        context = super().get_renderer_context()
        # Force the order from the serializer. Otherwise redundant literal list!
        # ser_cls = self.get_serializer_class()
        # instead of "ser_cls.Meta.fields if ser_cls else None":
        context["header"] = [
            "deployment.event_deployed_to.id",
            "deployment.event_deployed_to.glide",
            "deployment.event_deployed_to.name",
            "deployment.event_deployed_to.ifrc_severity_level",
            "deployment.event_deployed_to.dtype_name",
            "deployment.event_deployed_to.countries.name",
            "deployment.event_deployed_to.countries.iso3",
            "deployment.event_deployed_to.countries.society_name",
            "deployment.event_deployed_to.countries.region",
            "role",
            "type",
            "surge_alert_id",
            "appraisal_received",
            "gender",
            "location",
        ]
        context["header"] += [
            "id",
            "country_to.name",
            "country_to.iso3",
            "country_to.society_name",
            "country_to.region",
            "country_from.name",
            "country_from.iso3",
            "country_from.society_name",
            "country_from.region",
            "start_date",
            "end_date",
            "ongoing",
            "is_active",
            "name",
            "molnix_status",
        ]
        context["request"] = self.request
        context["header"] += [
            "molnix_id",
            "molnix_sector",
            "molnix_role_profile",
            "molnix_language",
            "molnix_region",
            "molnix_scope",
            "molnix_modality",
            "molnix_operation",
        ]
        if self.request.user.is_superuser:
            context["header"] += ["inactive_status"]

        context["labels"] = {i: i for i in context["header"]}
        # We can change the column titles (called "labels"):
        context["labels"]["deployment.event_deployed_to.id"] = "event_id"
        context["labels"]["deployment.event_deployed_to.glide"] = "event_glide_id"
        context["labels"]["deployment.event_deployed_to.name"] = "event_name"
        context["labels"]["deployment.event_deployed_to.ifrc_severity_level"] = "event_ifrc_severity_level"
        context["labels"]["deployment.event_deployed_to.dtype_name"] = "event_disaster_type"
        context["labels"]["deployment.event_deployed_to.countries.name"] = "event_country_name"
        context["labels"]["deployment.event_deployed_to.countries.iso3"] = "event_country_iso3"
        context["labels"]["deployment.event_deployed_to.countries.society_name"] = "event_country_nationalsociety"
        context["labels"]["deployment.event_deployed_to.countries.region"] = "event_country_regionname"
        context["labels"]["id"] = "deployed_id"
        context["labels"]["country_to.name"] = "deployed_to_name"
        context["labels"]["country_to.iso3"] = "deployed_to_iso3"
        context["labels"]["country_to.society_name"] = "deployed_to_nationalsociety"
        context["labels"]["country_to.region"] = "deployed_to_regionname"
        context["labels"]["country_from.name"] = "deployed_from_name"
        context["labels"]["country_from.iso3"] = "deployed_from_iso3"
        context["labels"]["country_from.society_name"] = "deployed_from_nationalsociety"
        context["labels"]["country_from.region"] = "deployed_from_regionname"
        context["labels"]["surge_alert_id"] = "surge_alert_id"
        context["labels"]["appraisal_received"] = "appraisal_received"
        context["labels"]["gender"] = "gender"
        context["labels"]["location"] = "location"

        # https://github.com/mjumbewu/django-rest-framework-csv/blob/master/rest_framework_csv/renderers.py
        # #L226-L229 uses bom when required:
        context["bom"] = True

        return context


class AggregateDeployments(APIView):
    """
    Get aggregated data for personnel deployments
    """

    @classmethod
    @extend_schema(request=None, responses=AggregateDeploymentsSerializer)
    def get(cls, request):
        today = timezone.now().date().strftime("%Y-%m-%d")
        this_year = timezone.now().year
        deployments_qset = Personnel.objects.filter(is_active=True)
        eru_qset = ERU.objects.all()
        if request.GET.get("event"):
            event_id = request.GET.get("event")
            deployments_qset = deployments_qset.filter(deployment__event_deployed_to=event_id)
            eru_qset = eru_qset.filter(event=event_id)

        active_rapid_response_personnel = deployments_qset.filter(
            type=Personnel.TypeChoices.RR,
            start_date__date__lte=today,
            end_date__date__gte=today,
            is_active=True,
        ).count()

        rapid_response_deployments_this_year = deployments_qset.filter(
            is_active=True, start_date__year__lte=this_year, end_date__year__gte=this_year
        ).count()
        active_emergency_response_units = eru_qset.filter(
            deployed_to__isnull=False,
            start_date__date__lte=today,
            end_date__date__gte=today,
        ).count()

        emergency_response_unit_deployed_this_year = eru_qset.filter(
            deployed_to__isnull=False,
            start_date__year__lte=this_year,
            end_date__year__gte=this_year,
        ).count()
        return Response(
            AggregateDeploymentsSerializer(
                dict(
                    active_rapid_response_personnel=active_rapid_response_personnel,
                    rapid_response_deployments_this_year=rapid_response_deployments_this_year,
                    active_emergency_response_units=active_emergency_response_units,
                    emergency_response_unit_deployed_this_year=emergency_response_unit_deployed_this_year,
                )
            ).data
        )


class DeploymentsByMonth(APIView):
    @classmethod
    @extend_schema(
        request=None,
        responses=DeploymentsByMonthSerializer(many=True),
    )
    def get(cls, request):
        """
        Returns count of Personnel Deployments
        for last 12 months, aggregated by month.
        """
        now = datetime.datetime.now()
        months = get_previous_months(now, 12)
        deployment_counts = []
        for month in months:
            month_string = month[0]
            first_day = month[1]
            last_day = month[2]
            count = Personnel.objects.filter(start_date__date__lte=last_day, end_date__date__gte=first_day).count()
            deployment_counts.append(dict(date=month_string, count=count))
        return Response(DeploymentsByMonthSerializer(deployment_counts, many=True).data)


class DeploymentsByNS(APIView):
    @classmethod
    @extend_schema(request=None, responses=DeploymentByNSSerializer(many=True))
    def get(cls, request):
        """Returns count of Personnel Deployments
        by National Society, for the current year.
        """
        limit = request.GET.get("limit", "5")
        limit = int(limit)
        now = datetime.datetime.now()
        first_day_of_year = date(now.year, 1, 1)
        last_day_of_year = date(now.year, 12, 31)
        societies = (
            Country.objects.filter(
                personnel_deployments__start_date__date__lte=last_day_of_year,
                personnel_deployments__end_date__date__gte=first_day_of_year,
            )
            .annotate(deployments_count=Count("personnel_deployments"))
            .order_by("-deployments_count")
            .values("id", "society_name", "deployments_count")[0:limit]
        )
        return Response(DeploymentByNSSerializer(societies, many=True).data)


class PartnerDeploymentFilterset(filters.FilterSet):
    parent_society = filters.NumberFilter(field_name="parent_society", lookup_expr="exact")
    country_deployed_to = filters.NumberFilter(field_name="country_deployed_to", lookup_expr="exact")
    district_deployed_to = filters.NumberFilter(field_name="district_deployed_to", lookup_expr="exact")
    parent_society__in = ListFilter(field_name="parent_society__id")
    country_deployed_to__in = ListFilter(field_name="country_deployed_to__id")
    district_deployed_to__in = ListFilter(field_name="district_deployed_to__id")

    class Meta:
        model = PartnerSocietyDeployment
        fields = {
            "start_date": ("exact", "gt", "gte", "lt", "lte"),
            "end_date": ("exact", "gt", "gte", "lt", "lte"),
        }


class PartnerDeploymentViewset(viewsets.ReadOnlyModelViewSet):
    queryset = PartnerSocietyDeployment.objects.all()
    serializer_class = PartnerDeploymentSerializer
    filterset_class = PartnerDeploymentFilterset

    def get_serializer_class(self):
        if is_tableau(self.request) is True:
            return PartnerDeploymentTableauSerializer
        return PartnerDeploymentSerializer


class RegionalProjectViewset(viewsets.ReadOnlyModelViewSet):
    queryset = RegionalProject.objects.all()
    serializer_class = RegionalProjectSerializer
    search_fields = ("name",)  # for /docs


class ProjectViewset(
    RevisionMixin,
    CsvListMixin,
    ReadOnlyVisibilityViewsetMixin,
    viewsets.ModelViewSet,
):
    filterset_class = ProjectFilter
    serializer_class = ProjectSerializer
    csv_serializer_class = ProjectCsvSerializer
    ordering_fields = ("name",)
    search_fields = ("name",)  # for /docs
    queryset = (
        Project.objects.select_related(
            "user", "modified_by", "project_country", "reporting_ns", "dtype", "regional_project", "primary_sector"
        )
        .prefetch_related(
            "project_districts",
            "event",
            "event__appeals",
            "event__countries_for_preview",
            "annual_splits",
            "secondary_sectors",
            "project_admin2",
        )
        .order_by("-modified_at")
        .all()
    )

    def get_permissions(self):
        # Require authentication for unsafe methods only
        if self.action in ["list", "retrieve"]:
            permission_classes = []
        else:
            permission_classes = [IsAuthenticated, DenyGuestUserPermission]
        return [permission() for permission in permission_classes]


class RegionProjectViewset(ReadOnlyVisibilityViewsetMixin, viewsets.ViewSet):
    def get_region(self):
        if not hasattr(self, "_region"):
            self._region = get_object_or_404(Region, pk=self.kwargs["pk"])
        return self._region

    def get_projects(self):
        # Region Filter will be applied using ProjectFilter if provided
        # Filter by visibility
        qs = self.get_visibility_queryset(Project.objects.all())
        if self.action != "global_national_society_activities":
            region = self.get_region()
            # Filter by region (From URL Params)
            qs = qs.filter(
                models.Q(project_country__region=region) | models.Q(project_districts__country__region=region)
            ).distinct()
        # Filter by GET params
        return ProjectFilter(self.request.query_params, queryset=qs).qs

    @extend_schema(request=None, responses=ProjectRegionOverviewSerializer)
    @action(detail=True, url_path="overview", methods=("get",))
    def overview(self, request, pk=None):
        projects = self.get_projects()
        aggregate_data = projects.aggregate(
            total_budget=models.Sum("budget_amount"),
            target_total=models.Sum("target_total"),
            reached_total=models.Sum("reached_total"),
        )
        data = {
            "total_projects": projects.count(),
            "ns_with_ongoing_activities": projects.filter(status=Statuses.ONGOING)
            .order_by("reporting_ns")
            .values("reporting_ns")
            .distinct()
            .count(),
            "total_budget": aggregate_data["total_budget"],
            "target_total": aggregate_data["target_total"],
            "reached_total": aggregate_data["reached_total"],
            "projects_by_status": projects.order_by()
            .values("status")
            .annotate(count=models.Count("id", distinct=True))
            .values("status", "count"),
        }
        return Response(data)

    @extend_schema(request=None, responses=ProjectRegionMovementActivitiesSerializer)
    @action(detail=True, url_path="movement-activities", methods=("get",))
    def movement_activities(self, request, pk=None):
        projects = self.get_projects()

        def _get_country_ns_sector_count():
            agg = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: defaultdict(int))))
            sectortitle = {}
            fields = (
                "project_country",
                "reporting_ns",
                "primary_sector",  # names below:
                "project_country__name",
                "reporting_ns__name",
                "primary_sector__title",
            )
            qs = (
                projects.order_by()
                .values(*fields[:3])
                .annotate(count=models.Count("id", distinct=True))
                .values_list(
                    *fields,
                    "count",
                )
            )
            for country, ns, sector, country_name, ns_name, title, count in qs:
                sectortitle[sector] = title
                agg[country][ns][sector] = count
                agg[country]["name"] = country_name
                agg[country][ns]["name"] = ns_name
            return [
                {
                    "id": cid,
                    "name": country.pop("name"),
                    "reporting_national_societies": [
                        {
                            "id": nsid,
                            "name": ns.pop("name"),
                            "sectors": [
                                {
                                    "id": sector,
                                    "sector": sectortitle[sector],
                                    "count": count,
                                }
                                for sector, count in ns.items()
                            ],
                        }
                        for nsid, ns in country.items()
                    ],
                }
                for cid, country in agg.items()
            ]

        region = self.get_region()
        country_projects = projects.filter(project_country=models.OuterRef("pk"))
        countries = Country.objects.filter(region=region)

        # Using english label for now
        with translation_override("en"):
            country_annotate = {
                f"{status_label.lower()}_projects_count": Coalesce(
                    models.Subquery(
                        country_projects.filter(status=status)
                        .values("project_country")
                        .annotate(count=models.Count("id", distinct=True))
                        .values("count")[:1],
                        output_field=models.IntegerField(),
                    ),
                    0,
                )
                for status, status_label in Statuses.choices
            }

        data = {
            "total_projects": projects.count(),
            "countries_count": countries.annotate(
                projects_count=Coalesce(
                    models.Subquery(
                        projects.filter(project_country=models.OuterRef("pk"))
                        .values("project_country")
                        .annotate(count=models.Count("id", distinct=True))
                        .values("count")[:1],
                        output_field=models.IntegerField(),
                    ),
                    0,
                ),
                **country_annotate,
            ).values("id", "name", "iso", "iso3", "projects_count", *country_annotate.keys()),
            "country_ns_sector_count": _get_country_ns_sector_count(),
            "supporting_ns": [
                {"id": id, "name": name, "count": count}
                for id, name, count in projects.order_by()
                .values("reporting_ns")
                .annotate(count=models.Count("id", distinct=True))
                .values_list("reporting_ns", "reporting_ns__name", "count")
            ],
        }
        return Response(ProjectRegionMovementActivitiesSerializer(data).data)

    @action(detail=True, url_path="national-society-activities", methods=("get",))
    def national_society_activities(self, request, pk=None):
        projects = self.get_projects()
        title = {t.id: t.title for t in Sector.objects.all()}

        def _get_distinct(field, *args, **kwargs):
            kwargs[field] = field
            return [
                {f: p[key] for f, key in kwargs.items()}
                for p in projects.order_by()
                .values(field)
                .annotate(count=models.Count("id", distinct=True))
                .values(field, *kwargs.values())
                .distinct()
            ]

        def _get_count(*fields):
            return list(
                projects.order_by()
                .values(*fields)
                .annotate(count=models.Count("id", distinct=True))
                .values_list(*fields, "count")
            )

        # Raw nodes
        supporting_ns_list = _get_distinct(
            "reporting_ns",
            iso3="reporting_ns__iso3",
            iso="reporting_ns__iso",
            name="reporting_ns__society_name",
        )
        receiving_ns_list = _get_distinct(
            "project_country",
            iso3="project_country__iso3",
            iso="project_country__iso",
            name="project_country__name",
        )
        sector_list = _get_distinct("primary_sector")

        # Raw links
        supporting_ns_and_sector_group = _get_count("reporting_ns", "primary_sector")
        sector_and_receiving_ns_group = _get_count("primary_sector", "project_country")

        # Node Types
        SUPPORTING_NS = "supporting_ns"
        RECEIVING_NS = "receiving_ns"
        SECTOR = "sector"

        nodes = [
            {
                "id": node[id_selector],
                "type": gtype,
                **(
                    {
                        "name": node["name"],
                        "iso": node["iso"],
                        "iso3": node["iso3"],
                    }
                    if gtype != SECTOR
                    else {
                        "name": title[node[id_selector]],
                    }
                ),
            }
            for group, gtype, id_selector in [
                (supporting_ns_list, SUPPORTING_NS, "reporting_ns"),
                (sector_list, SECTOR, "primary_sector"),
                (receiving_ns_list, RECEIVING_NS, "project_country"),
            ]
            for node in group
        ]

        node_id_map = {f"{node['type']}-{node['id']}": index for index, node in enumerate(nodes)}

        links = [
            {
                "source": node_id_map[f"{source_type}-{source}"],
                "target": node_id_map[f"{target_type}-{target}"],
                "value": value,
            }
            for group, source_type, target_type in [
                (supporting_ns_and_sector_group, SUPPORTING_NS, SECTOR),
                (sector_and_receiving_ns_group, SECTOR, RECEIVING_NS),
            ]
            for source, target, value in group
        ]
        return Response({"nodes": nodes, "links": links})

    @action(detail=False, url_path="national-society-activities", methods=("get",))
    def global_national_society_activities(self, request, pk=None):
        return self.national_society_activities(request, pk)


class GlobalProjectViewset(ReadOnlyVisibilityViewsetMixin, viewsets.ViewSet):
    def get_projects(self):
        # Filter by visibility
        qs = self.get_visibility_queryset(Project.objects.all())
        # Filter by GET params
        projects = ProjectFilter(self.request.query_params, queryset=qs).qs
        return Project.objects.filter(
            # To avoid duplicate rows
            id__in=projects,
            # NOTE: Only process ongoing projects in global project view
            status=Statuses.ONGOING,
        )

    @extend_schema(methods=["GET"], request=None, responses=GlobalProjectOverviewSerializer)
    @action(detail=False, url_path="overview", methods=("get",))
    def overview(self, request, pk=None):
        def _get_projects_per_enum_field(EnumType, enum_field):
            """
            Use this for enum fields
            """
            return [
                {
                    enum_field: enum_field_value,
                    f"{enum_field}_display": EnumType(int(enum_field_value)).label,
                    "count": count,
                }
                for enum_field_value, count in (
                    projects.order_by()
                    .values(enum_field)
                    .annotate(count=models.Count("id"))
                    .values_list(
                        enum_field,
                        "count",
                    )
                    .order_by(enum_field)
                )
            ]

        def _get_projects_per_foreign_field(field, field_display):
            """
            Use this for foreign fields
            """
            return [
                {
                    field: field_value,
                    f"{field}_display": field_display_value,
                    "count": count,
                }
                for field_value, field_display_value, count in (
                    projects.order_by()
                    .values(field, field_display)
                    .annotate(count=models.Count("id"))
                    .values_list(
                        field,
                        field_display,
                        "count",
                    )
                    .order_by(field)
                )
                if field_value is not None
            ]

        projects = self.get_projects()

        target_total = projects.aggregate(target_total=models.Sum("target_total"))["target_total"]
        response_data = {
            "total_ongoing_projects": projects.filter(status=Statuses.ONGOING).count(),
            "ns_with_ongoing_activities": (
                projects.filter(status=Statuses.ONGOING).order_by("reporting_ns").values("reporting_ns").distinct().count()
            ),
            "target_total": target_total,
            "projects_per_sector": _get_projects_per_foreign_field("primary_sector", "primary_sector__title"),
            "projects_per_programme_type": _get_projects_per_enum_field(ProgrammeTypes, "programme_type"),
            "projects_per_secondary_sectors": _get_projects_per_foreign_field("secondary_sectors", "secondary_sectors__title"),
        }
        return Response(GlobalProjectOverviewSerializer(response_data).data)

    @extend_schema(methods=["GET"], request=None, responses=GlobalProjectNSOngoingProjectsStatsSerializer(many=True))
    @action(detail=False, url_path="ns-ongoing-projects-stats", methods=("get",))
    def ns_ongoing_projects_stats(self, request, pk=None):
        projects = self.get_projects()
        ref_projects = projects.filter(reporting_ns=models.OuterRef("pk"))

        project_per_sector = defaultdict(list)
        for reporting_ns, primary_sector, primary_sector__title, count in (
            projects.order_by("reporting_ns", "primary_sector")
            .values("reporting_ns", "primary_sector")
            .annotate(count=models.Count("id"))
            .values_list("reporting_ns", "primary_sector", "primary_sector__title", "count")
        ):
            project_per_sector[reporting_ns].append(
                {
                    "primary_sector": primary_sector,
                    "primary_sector_display": primary_sector__title,
                    "count": count,
                }
            )

        response_data = [
            {
                **ns_data,
                "projects_per_sector": project_per_sector.get(ns_data["id"]),
                "operation_types_display": [
                    OperationTypes(operation_type).label for operation_type in ns_data["operation_types"]
                ],
            }
            for ns_data in Country.objects.annotate(
                ongoing_projects=Coalesce(
                    models.Subquery(
                        ref_projects.values("reporting_ns").annotate(count=models.Count("id")).values("count")[:1],
                        output_field=models.IntegerField(),
                    ),
                    0,
                ),
                target_total=Coalesce(
                    models.Subquery(
                        ref_projects.values("reporting_ns")
                        .annotate(target_total=models.Sum("target_total"))
                        .values("target_total")[:1],
                        output_field=models.IntegerField(),
                    ),
                    0,
                ),
                budget_amount_total=Coalesce(
                    models.Subquery(
                        ref_projects.values("reporting_ns")
                        .annotate(budget_amount_total=models.Sum("budget_amount"))
                        .values("budget_amount_total")[:1],
                        output_field=models.IntegerField(),
                    ),
                    0,
                ),
                operation_types=Coalesce(
                    models.Subquery(
                        ref_projects.values("reporting_ns")
                        .annotate(operation_types=ArrayAgg("operation_type", distinct=True))
                        .values("operation_types")[:1],
                        output_field=ArrayField(models.IntegerField()),
                    ),
                    [],
                ),
            )
            .filter(ongoing_projects__gt=0)
            .order_by("id")
            .values(
                "id",
                "name",
                "iso3",
                "society_name",
                "ongoing_projects",
                "target_total",
                "budget_amount_total",
                "operation_types",
            )
        ]
        return Response(GlobalProjectNSOngoingProjectsStatsSerializer(response_data, many=True).data)


class EmergencyProjectViewSet(
    RevisionMixin,
    # ReadOnlyVisibilityViewsetMixin,  # FIXME: This is required?
    viewsets.ModelViewSet,
):
    queryset = (
        EmergencyProject.objects.select_related("created_by", "reporting_ns", "event", "country", "deployed_eru", "modified_by")
        .prefetch_related(
            "districts",
            "admin2",
            "event__appeals",
            "event__countries_for_preview",
            "activities",
            "activities__sector",
            "activities__action",
            "activities__action__supplies",
            "activities__points",
        )
        .order_by("-modified_at")
        .all()
    )
    # Intentionally not IsAuthenticated. Anons should see public EmergencyProjects:
    permission_classes = []
    filterset_class = EmergencyProjectFilter
    serializer_class = EmergencyProjectSerializer
    ordering_fields = ("title",)
    search_fields = ("title",)  # for /docs

    @action(
        detail=False,
        url_path="options",
        methods=("get",),
        serializer_class=EmergencyProjectOptionsSerializer,
    )
    def get_options(self, request, pk=None):
        return Response(
            EmergencyProjectOptionsSerializer(
                instance=dict(
                    sectors=EmergencyProjectActivitySector.objects.all(),
                    actions=EmergencyProjectActivityAction.objects.prefetch_related("supplies").all(),
                )
            ).data
        )


class ERUReadinessFilter(filters.FilterSet):
    eru_owner = filters.NumberFilter(field_name="eru_owner", lookup_expr="exact")
    eru_type = filters.NumberFilter(field_name="eru_types__type", lookup_expr="exact")
    eru_type__in = ListFilter(field_name="eru_types__type")


class ERUReadinessViewSet(RevisionMixin, viewsets.ModelViewSet):
    queryset = ERUReadiness.objects.all()
    serializer_class = ERUReadinessSerializer
    filterset_class = ERUReadinessFilter
    permission_classes = [ERUReadinessPermission]
    ordering_fields = "__all__"

    def get_queryset(self):
        return (
            super()
            .get_queryset()
            .select_related(
                "eru_owner__national_society_country",
            )
            .prefetch_related(
                "eru_types",
            )
            .order_by("-updated_at")
        )

    def delete(self, request, *args, **kwargs):
        return bad_request("Delete method not allowed")


class ERUReadinessTypeFilter(filters.FilterSet):
    eru_owner = filters.NumberFilter(field_name="erureadiness__eru_owner", lookup_expr="exact", label="ERU Owner")
    type = filters.NumberFilter(field_name="type", lookup_expr="exact")


class ERUReadinessTypeViewset(viewsets.ReadOnlyModelViewSet):
    queryset = ERUReadinessType.objects.all()
    serializer_class = MiniERUReadinessTypeSerializer
    filterset_class = ERUReadinessTypeFilter
    ordering_fields = "__all__"

    def get_queryset(self):
        qs = super().get_queryset()
        return (
            qs.filter(
                erureadiness__isnull=False,
            )
            .prefetch_related(
                "erureadiness_set__eru_owner__national_society_country",
            )
            .distinct()
        )


class ExportERUReadinessView(APIView):
    def get(self, request, *args, **kwargs):
        wb = Workbook()
        ws = wb.active
        ws.title = "ERU Readiness"

        static_headers = ["National Society", "Updated Date"]
        main_headers = static_headers.copy()  # Create a separate list to keep static_headers unchanged
        sub_headers = [""] * len(static_headers)
        readiness_columns = ["Equipment", "People", "Funding", "Comment"]

        for type_label in ERUType.labels:
            main_headers.append(str(type_label))
            main_headers.extend(
                [""] * (len(readiness_columns) - 1)
            )  # Fill empty cells to align merged ERU type header across subcolumns

            sub_headers.extend(readiness_columns)

        ws.append(main_headers)
        ws.append(sub_headers)

        column_start = len(static_headers) + 1  # Determine starting column for merging
        for _ in ERUType.choices:
            column_end = column_start + len(readiness_columns) - 1
            ws.merge_cells(start_row=1, start_column=column_start, end_row=1, end_column=column_end)
            column_start += len(readiness_columns)

        eru_readiness_queryset = ERUReadiness.objects.select_related(
            "eru_owner__national_society_country",
        ).prefetch_related("eru_types")

        for eru_readiness in eru_readiness_queryset.iterator():
            row_data = [
                eru_readiness.eru_owner.national_society_country.name,
                eru_readiness.updated_at.strftime("%Y-%m-%d"),
            ]

            readiness_data_mapping = {
                eru_readiness_type.type: {
                    "equipment": eru_readiness_type.get_equipment_readiness_display(),
                    "people": eru_readiness_type.get_people_readiness_display(),
                    "funding": eru_readiness_type.get_funding_readiness_display(),
                    "comment": eru_readiness_type.comment if eru_readiness_type.comment else "",
                }
                for eru_readiness_type in eru_readiness.eru_types.all()
            }

            for eru_type_value in ERUType.values:
                if eru_type_value in readiness_data_mapping:
                    row_data.extend(readiness_data_mapping[eru_type_value].values())
                else:
                    row_data.extend(["" for _ in readiness_columns])  # Empty placeholders

            ws.append(row_data)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = "attachment; filename=eru_readiness_export.xlsx"
        wb.save(response)
        return response
